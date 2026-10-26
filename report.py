"""
report.py
---------
Exportación de un reporte PDF de resultados (tabla de indicadores +
gráficos) para el conjunto de datos analizado en la interfaz gráfica.
Este PDF es distinto del PDF teórico de la herramienta: aquí se
genera un reporte específico de la corrida/análisis que el usuario
haya hecho con sus propios datos, con fecha y parámetros usados.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path
from datetime import datetime

import pandas as pd
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib.utils import ImageReader
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_JUSTIFY
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                 TableStyle, Image, PageBreak, KeepTogether)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from .stats import INDICATOR_HELP, INDICATOR_SCALE, quality_label, scale_description, calculation_trace
from .plots import FIGURE_DESCRIPTIONS, indicator_gauge_plot

NAVY_HEX = "2C3E50"
APP_NAME = "CT · Comparison Tool"
AUTHOR_NAME = "PhD. Mauricio Villazón"
AUTHOR_LINE = (f"Autor de la metodología: {AUTHOR_NAME} — creador de la herramienta CT y de la "
               f"planilla original CT_2010, basada en los trabajos de Vázquez, Willems, "
               f"Abuelnasr y Biesbrouck (K.U.Leuven) y el prof. Bauwens (VUB).")

# ---------------------------------------------------------------------------
# Estilo de reporte configurable por la usuaria (tipografía, alineación de
# párrafos y estilo de tabla): elegido en la interfaz (diálogo "Estilo de
# reporte") y pasado como un único diccionario `style` a los tres
# generadores (`build_results_pdf`/`_docx`/`_xlsx`). Si no se pasa nada, se
# usan estos valores por defecto, que reproducen el aspecto que tenía el
# reporte antes de esta función (ningún cambio visual para quien no toque
# el diálogo nuevo).
# ---------------------------------------------------------------------------
REPORT_FONT_CHOICES = ["Arial", "Times New Roman", "Calibri", "Cambria"]
REPORT_ALIGN_CHOICES = ["left", "justify", "center"]
REPORT_TABLE_STYLE_CHOICES = ["moderno", "formal"]
DEFAULT_REPORT_STYLE = {
    "font_family": "Arial",
    "paragraph_align": "left",
    "table_style": "moderno",
}

# Fuentes TrueType métricamente compatibles disponibles en el sistema, que
# se registran en reportlab bajo el nombre "de marca" que ve la usuaria
# (Arial/Times New Roman/Calibri/Cambria), para que el PDF cambie de
# tipografía de verdad y no solo de etiqueta. Si algún archivo no existe
# (por ejemplo, en otro sistema operativo sin estos paquetes), se cae en
# silencio a las fuentes base de reportlab (Helvetica), sin romper la
# generación del reporte.
_FONT_FILES = {
    "Arial": {
        "normal": "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "bold": "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
        "italic": "/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf",
        "boldItalic": "/usr/share/fonts/truetype/liberation/LiberationSans-BoldItalic.ttf",
    },
    "Times New Roman": {
        "normal": "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf",
        "bold": "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf",
        "italic": "/usr/share/fonts/truetype/liberation/LiberationSerif-Italic.ttf",
        "boldItalic": "/usr/share/fonts/truetype/liberation/LiberationSerif-BoldItalic.ttf",
    },
    "Calibri": {
        "normal": "/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf",
        "bold": "/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf",
        "italic": "/usr/share/fonts/truetype/crosextra/Carlito-Italic.ttf",
        "boldItalic": "/usr/share/fonts/truetype/crosextra/Carlito-BoldItalic.ttf",
    },
    "Cambria": {
        "normal": "/usr/share/fonts/truetype/crosextra/Caladea-Regular.ttf",
        "bold": "/usr/share/fonts/truetype/crosextra/Caladea-Bold.ttf",
        "italic": "/usr/share/fonts/truetype/crosextra/Caladea-Italic.ttf",
        "boldItalic": "/usr/share/fonts/truetype/crosextra/Caladea-BoldItalic.ttf",
    },
}
_FONT_FAMILIES_REGISTERED: set[str] = set()


def _register_font_family(family: str) -> bool:
    """Registra en reportlab la familia tipográfica `family` (si todavía
    no se registró) a partir de sus archivos .ttf. Devuelve True si quedó
    disponible (ya sea porque se registró ahora o en una llamada previa),
    False si no se pudo (archivo faltante u otro error), en cuyo caso
    quien llama debe usar Helvetica como respaldo."""
    if family in _FONT_FAMILIES_REGISTERED:
        return True
    files = _FONT_FILES.get(family)
    if not files or not all(Path(p).exists() for p in files.values()):
        return False
    try:
        safe = family.replace(" ", "")
        name_n, name_b = safe, f"{safe}-Bold"
        name_i, name_bi = f"{safe}-Italic", f"{safe}-BoldItalic"
        pdfmetrics.registerFont(TTFont(name_n, files["normal"]))
        pdfmetrics.registerFont(TTFont(name_b, files["bold"]))
        pdfmetrics.registerFont(TTFont(name_i, files["italic"]))
        pdfmetrics.registerFont(TTFont(name_bi, files["boldItalic"]))
        pdfmetrics.registerFontFamily(safe, normal=name_n, bold=name_b,
                                       italic=name_i, boldItalic=name_bi)
        _FONT_FAMILIES_REGISTERED.add(family)
        return True
    except Exception:
        return False


def _pdf_font_names(family: str) -> dict:
    """Nombres de fuente reportlab (normal/bold/italic/boldItalic) para la
    familia elegida por la usuaria, con Helvetica como respaldo seguro."""
    if _register_font_family(family):
        safe = family.replace(" ", "")
        return dict(normal=safe, bold=f"{safe}-Bold",
                    italic=f"{safe}-Italic", boldItalic=f"{safe}-BoldItalic")
    return dict(normal="Helvetica", bold="Helvetica-Bold",
                italic="Helvetica-Oblique", boldItalic="Helvetica-BoldOblique")


_PDF_ALIGN = {"left": TA_LEFT, "justify": TA_JUSTIFY, "center": TA_CENTER}


def _normalize_style(style: dict | None) -> dict:
    out = dict(DEFAULT_REPORT_STYLE)
    if style:
        for k in ("font_family", "paragraph_align", "table_style"):
            if style.get(k):
                out[k] = style[k]
    return out


def _generic_table_style(table_style: str, font_bold: str, font_normal: str,
                          zebra: bool = True, font_size: float = 8) -> TableStyle:
    """Comandos de estilo compartidos por las tablas "sueltas" del reporte
    (parámetros, resultado por indicador, ranking...), con dos variantes:
    "moderno" (encabezado azul marino, filas alternadas en gris claro -- el
    aspecto original) y "formal" (encabezado blanco con texto en negro y
    SOLO tres filetes horizontales -- arriba del encabezado, bajo el
    encabezado y al pie de la tabla -- sin grilla ni líneas verticales ni
    entre filas de datos: estilo "booktabs" de tabla de artículo
    científico, como las tablas de materiales de presas CFRD de la
    bibliografía que sirvieron de referencia)."""
    if table_style == "formal":
        cmds = [
            ("FONTNAME", (0, 0), (-1, 0), font_bold),
            ("FONTNAME", (0, 1), (-1, -1), font_normal),
            ("BACKGROUND", (0, 0), (-1, 0), colors.white),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
            ("LINEABOVE", (0, 0), (-1, 0), 1.1, colors.black),
            ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.black),
            ("LINEBELOW", (0, -1), (-1, -1), 1.1, colors.black),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("TOPPADDING", (0, 0), (-1, -1), 3.2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3.2),
            ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ]
    else:
        cmds = [
            ("FONTNAME", (0, 0), (-1, 0), font_bold),
            ("FONTNAME", (0, 1), (-1, -1), font_normal),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2c3e50")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ]
        if zebra:
            cmds.append(("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white]))
    return TableStyle(cmds)


def _base_dir() -> Path:
    """Misma lógica que `gui_qt.main_window._base_dir`: carpeta base del
    proyecto, tanto ejecutando desde el código fuente como desde un
    .exe compilado con PyInstaller (--onefile extrae los datos
    empaquetados a una carpeta temporal apuntada por sys._MEIPASS).
    Se repite aquí (en vez de importar `gui_qt`) porque `report.py` no
    debe depender de PyQt: tiene que poder usarse sin la GUI."""
    if getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS"):
        return Path(sys._MEIPASS)
    return Path(__file__).resolve().parent.parent


WATERMARK_IMAGE_PATH = _base_dir() / "assets" / "villazon_logo.png"

# Colores de fondo suaves para las filas de resultados de cada
# indicador en el reporte, según su calificación (mismo criterio que
# `theme.kpi_color_for`/`stats.quality_label` en la interfaz).
_QUALITY_BG_HEX = {"Bueno": "E3F6EC", "Regular": "FCF3DC", "Malo": "FBE7E7"}

# ---------------------------------------------------------------------------
# Tabla resumen de indicadores (PDF): anchos de columna PROPORCIONALES
# (ver `_summary_col_widths_cm`) que siempre suman exactamente el ancho
# útil de una hoja A4 vertical (21cm - 2×1.8cm de margen = 17.4cm), sin
# importar cuántas columnas termine teniendo la tabla -- desde la
# versión que agrega el selector de qué indicadores incluir en el
# reporte, el número de columnas ya no es fijo (la usuaria puede
# des-marcar indicadores), así que un ancho fijo por columna podía
# volver a sumar más que la hoja y recortar columnas fuera del margen
# (bug ya corregido antes con anchos fijos, cuando la tabla era siempre
# de 16 columnas). Los pesos de abajo son relativos (no cm): una
# columna de peso 1.45 ocupa el doble de ancho que una de peso 0.725,
# sea cual sea el total de columnas presentes. Encabezados largos se
# parten en dos líneas para verse bien incluso angostos.
# ---------------------------------------------------------------------------
_SUMMARY_COL_WIDTHS_CM = {
    "Corrida": 2.2, "n": 0.7, "Avg(Oi)": 1.1, "Avg(Pi)": 1.1,
    "MAE": 0.85, "STD": 0.85, "RMSE": 0.85, "EF": 0.8, "CD": 0.8, "R": 0.7,
    "Model Skill (Willmott)": 1.45, "CR1 (Sesgo del modelo)": 1.45,
    "CR4 (NS logarítmico)": 1.35, "CR5 (NS adaptado)": 1.35,
    "PBIAS (%)": 0.95, "KGE": 0.8,
    "RSR": 0.75, "MAE/STD(Oi)": 1.2,
    "SEE (Auscultación de presas)": 1.55, "% dentro de ±2·SEE": 1.35,
    "Puntaje compuesto (0-1)": 1.35,
}
_DEFAULT_COL_WEIGHT = 1.0
_SUMMARY_HEADER_BREAKS = {
    "Avg(Oi)": "Avg<br/>(Oi)",
    "Avg(Pi)": "Avg<br/>(Pi)",
    "Model Skill (Willmott)": "Model Skill<br/>(Willmott)",
    "CR1 (Sesgo del modelo)": "CR1<br/>(Sesgo del modelo)",
    "CR4 (NS logarítmico)": "CR4<br/>(NS logarítmico)",
    "CR5 (NS adaptado)": "CR5<br/>(NS adaptado)",
    "PBIAS (%)": "PBIAS<br/>(%)",
    "MAE/STD(Oi)": "MAE/STD<br/>(Oi)",
    "SEE (Auscultación de presas)": "SEE<br/>(Auscultación)",
    "% dentro de ±2·SEE": "%<br/>dentro ±2·SEE",
}


def _summary_col_widths_cm(header: list[str], available_cm: float = 17.4) -> list[float]:
    """Reparte `available_cm` entre las columnas de `header` en
    proporción a `_SUMMARY_COL_WIDTHS_CM` (o un peso por defecto para
    columnas sin peso propio), de forma que la suma sea SIEMPRE exactamente
    `available_cm`, sin importar cuántas columnas haya (la tabla nunca
    se desborda del ancho de página, se achique o agrande el número de
    indicadores elegidos)."""
    weights = [_SUMMARY_COL_WIDTHS_CM.get(h, _DEFAULT_COL_WEIGHT) for h in header]
    total_weight = sum(weights) or 1.0
    return [w / total_weight * available_cm for w in weights]


def _build_summary_table(df: pd.DataFrame, header: list[str],
                          font_names: dict | None = None,
                          table_style: str = "moderno") -> Table:
    """Tabla de indicadores con encabezados en dos líneas y anchos de
    columna proporcionales (ver `_summary_col_widths_cm`), para que
    cualquier subconjunto de columnas elegido quepa completo en la
    misma hoja A4 sin recortarse, y celdas envueltas en `Paragraph`
    (no texto plano) para que un nombre de corrida largo también se
    ajuste a su columna en vez de desbordarla.

    `table_style`: "moderno" (encabezado azul marino, filas alternadas --
    aspecto original) o "formal" (encabezado blanco con texto en negro y
    doble filete horizontal, sin filas alternadas -- estilo de tabla
    académica/de publicación, como las tablas de materiales de presas CFRD
    en la bibliografía)."""
    font_names = font_names or dict(normal="Helvetica", bold="Helvetica-Bold")
    formal = table_style == "formal"
    head_color = colors.black if formal else colors.white
    head_style = ParagraphStyle("SummaryHeadCT", fontName=font_names["bold"], fontSize=6.3,
                                 leading=7.4, textColor=head_color, alignment=1)
    body_style = ParagraphStyle("SummaryBodyCT", fontName=font_names["normal"], fontSize=6.5,
                                 leading=7.6, alignment=1)
    body_left_style = ParagraphStyle("SummaryBodyLeftCT", parent=body_style, alignment=0)

    col_widths = [w * cm for w in _summary_col_widths_cm(header)]

    header_row = [Paragraph(_SUMMARY_HEADER_BREAKS.get(h, h), head_style) for h in header]
    body_rows = []
    for _, row in df.iterrows():
        cells = []
        for j, h in enumerate(header):
            val = row[h]
            text = f"{val:.4g}" if isinstance(val, float) else str(val)
            style = body_left_style if j == 0 else body_style
            cells.append(Paragraph(text, style))
        body_rows.append(cells)

    t = Table([header_row] + body_rows, hAlign="LEFT", colWidths=col_widths, repeatRows=1)
    if formal:
        style_cmds = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.white),
            ("LINEABOVE", (0, 0), (-1, 0), 1.1, colors.black),
            ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.black),
            ("LINEBELOW", (0, -1), (-1, -1), 1.1, colors.black),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.8),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]
    else:
        style_cmds = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2c3e50")),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
        ]
    t.setStyle(TableStyle(style_cmds))
    return t


def _indicator_results_rows(df: pd.DataFrame, key: str):
    """Lista de (nombre_corrida, valor, calificación) para un
    indicador, a partir de la tabla de resultados ya con "Corrida"
    como columna (formato que reciben los tres builders). Se comparte
    entre PDF/Word/Excel para no repetir esta lógica tres veces."""
    if key not in df.columns:
        return []
    name_col = "Corrida" if "Corrida" in df.columns else df.columns[0]
    out = []
    for _, row in df.iterrows():
        val = row[key]
        out.append((str(row[name_col]), val, quality_label(key, val)))
    return out


def _indicator_gauge_figure(df: pd.DataFrame, key: str):
    """Figura del medidor/escala de colores de un indicador (la misma
    que se ve en la pestaña Validación de la GUI), para incrustar en
    el reporte junto a la teoría/fórmula/pasos -- responde a si estos
    métodos se pueden graficar: sí, y no solo con texto. Devuelve
    `None` en vez de propagar la excepción si no se puede calcular
    (por ejemplo, muy pocos datos) — es un gráfico de apoyo, la
    descripción textual de la escala ya cubre esa información."""
    rows = _indicator_results_rows(df, key)
    if not rows:
        return None
    values_by_run = {name: val for name, val, _ in rows if isinstance(val, (int, float))}
    if not values_by_run:
        return None
    scale = INDICATOR_SCALE.get(key, dict(gauge_type="open_error"))
    obs_mean = float(df["Avg(Oi)"].iloc[0]) if "Avg(Oi)" in df.columns else None
    try:
        return indicator_gauge_plot(key, values_by_run, scale, obs_mean=obs_mean)
    except Exception:
        return None


def _calc_trace_lines(df: pd.DataFrame, key: str,
                       calc_detail: dict[str, dict] | None) -> list[str]:
    """Una línea de texto (ya en HTML simple) por corrida con "el paso
    del cálculo" de un indicador: con qué datos se trabajó, los
    cálculos intermedios reales (no solo descritos en teoría) y el
    resultado final -- reutiliza `ct_stats.calculation_trace` sobre
    los valores intermedios que ya trae `calc_detail` (un diccionario
    {nombre_corrida: detalle}, tal como lo arma la interfaz al
    ejecutar el análisis). Devuelve [] si no se pasó `calc_detail`
    (p. ej. quien llama a este módulo sin pasar por la GUI)."""
    if not calc_detail:
        return []
    rows = _indicator_results_rows(df, key)
    lines = []
    for run_name, val, _ in rows:
        detalle = calc_detail.get(run_name)
        if not detalle:
            continue
        try:
            trace = calculation_trace(key, detalle, val)
        except Exception:
            continue
        parts = [f"{label} = {value}" for label, value in trace]
        lines.append(f"<b>{run_name}:</b> " + " &nbsp;·&nbsp; ".join(parts))
    return lines


# ---------------------------------------------------------------------------
# Texto metodológico general (didáctico): es siempre el mismo, parafraseado
# de la teoría de la herramienta CT (ver CT_2010_explicacion_teorica.pdf);
# lo único que cambia entre reportes son los datos, parámetros y resultados
# de cada corrida. Se usa en las tres salidas (PDF/Word/Excel) para que el
# reporte sea más completo y autoexplicativo, sin depender de que quien lo
# lea también tenga el PDF teórico a mano.
# ---------------------------------------------------------------------------
METHODOLOGY_INTRO = (
    "Este reporte resume la comparación entre una serie observada y una o más "
    "series simuladas de caudal, siguiendo la metodología de la herramienta CT "
    "(Comparison Tool), basada en los trabajos de Vázquez (2003), Willmott "
    "(1981), Nash y Sutcliffe (1970) y Willems (2000, 2009), entre otros. El "
    "objetivo es cuantificar, con varios indicadores complementarios, qué tan "
    "bien reproduce el modelo el comportamiento real de la cuenca, tanto en su "
    "tendencia general como en sus eventos extremos (crecidas y estiajes)."
)
METHODOLOGY_INDICADORES = (
    "Los indicadores de bondad de ajuste se calculan sobre toda la serie de "
    "tiempo, comparando cada valor simulado (Pi) con el observado (Oi) en el "
    "mismo instante. Ningún indicador por sí solo cuenta toda la historia: "
    "MAE, STD y RMSE miden el error en las unidades originales del caudal; EF, "
    "CD, R, el índice de Willmott y KGE son adimensionales y comparan el "
    "modelo contra referencias (el promedio observado, la correlación, la "
    "variabilidad); PBIAS y CR1 miden específicamente el sesgo de volumen; "
    "CR4/CR5 dan más peso a los caudales bajos o a ambos extremos, para no "
    "depender solo de qué tan bien se ajustan las crecidas; y RSR y "
    "MAE/STD(Oi) normalizan el RMSE y el MAE por la variabilidad natural de "
    "la serie observada (Moriasi et al., 2007; Singh et al., 2004), lo que "
    "los hace comparables entre variables y estudios distintos."
)
METHODOLOGY_AUSCULTACION = (
    "Adicionalmente se incluye el SEE (Standard Error of Estimate), el "
    "indicador central de los modelos de auscultación de presas de tipo HST "
    "(Salazar et al., 2017), usado en la práctica para el monitoreo de "
    "instrumentación geotécnica (piezómetros, extensómetros, péndulos, "
    "etc.): se compara cada lectura contra la predicción del modelo ± un "
    "múltiplo del SEE, y las lecturas fuera de esa banda (típicamente "
    "±2·SEE) se marcan como candidatas a revisión, ya que pueden indicar un "
    "cambio real de comportamiento de la estructura. El porcentaje de "
    "lecturas dentro de ±2·SEE es el criterio de \"comportamiento "
    "satisfactorio\" que acompaña a este indicador en el reporte."
)
METHODOLOGY_EXTREMOS = (
    "Además del ajuste general, se analiza el comportamiento en los eventos "
    "extremos: se seleccionan automáticamente los picos de crecida y los "
    "mínimos de estiaje independientes entre sí (con una separación temporal "
    "mínima configurable), se calcula su periodo de retorno empírico, y se "
    "compara la serie de eventos observados contra la simulada. Para "
    "estabilizar la varianza de estos valores extremos —muy asimétricos por "
    "naturaleza— se aplica la transformación de Box-Cox (Box y Cox, 1964) a "
    "los máximos, y a los recíprocos (1/caudal) en el caso de los mínimos; el "
    "parámetro lambda de la transformación se estima automáticamente por "
    "máxima verosimilitud."
)


def _ranking_verdict_text(ranking_df: pd.DataFrame, used: list[str]) -> str:
    best_run = str(ranking_df.index[0])
    best_score = float(ranking_df["Puntaje compuesto (0-1)"].iloc[0])
    return (
        f"Para resumir en un solo número qué tan bien se ajusta cada corrida, "
        f"se calculó un puntaje compuesto (0 a 1, donde 1 es el ajuste "
        f"perfecto) promediando la calidad de {len(used)} indicadores clave "
        f"({', '.join(used)}), cada uno normalizado a esa misma escala. Este "
        f"puntaje es una referencia visual rápida —no reemplaza revisar cada "
        f"indicador por separado— y en esta corrida la mejor serie según ese "
        f"criterio combinado fue \"{best_run}\", con un puntaje de "
        f"{best_score:.3f}."
    )


def _multi_ranking_text(ranking_df: pd.DataFrame, used: list[str]) -> str:
    best_model = str(ranking_df.index[0])
    best_score = float(ranking_df["Puntaje promedio (0-1)"].iloc[0])
    n_vars = int(ranking_df["N° de variables"].iloc[0]) if "N° de variables" in ranking_df.columns else None
    vars_txt = f" (calculado sobre {n_vars} variable(s))" if n_vars else ""
    return (
        f"Del módulo de análisis múltiple/en paralelo: para cada variable analizada "
        f"(caudal, piezómetros, u otras) se calculó el mismo puntaje compuesto de "
        f"{len(used)} indicadores clave que se usa en el resto del reporte, y luego se "
        f"promedió ese puntaje por modelo a través de todas las variables, para resumir "
        f"en un solo número qué tan bien se ajusta cada modelo en conjunto a toda la "
        f"validación —no solo a una variable a la vez—. El modelo con mejor puntaje "
        f"promedio fue \"{best_model}\", con {best_score:.3f}{vars_txt}."
    )


def _build_multi_ranking_table(ranking_df: pd.DataFrame, font_names: dict | None = None,
                                table_style: str = "moderno") -> Table:
    font_names = font_names or dict(normal="Helvetica", bold="Helvetica-Bold")
    cols = [c for c in ["Puntaje promedio (0-1)", "N° de variables"] if c in ranking_df.columns]
    rk = ranking_df[cols].reset_index()
    rk.columns = ["Modelo"] + cols
    rows = [list(rk.columns)] + rk.round(4).astype(str).values.tolist()
    t = Table(rows, hAlign="LEFT", colWidths=[6 * cm] + [4.5 * cm] * len(cols))
    cmds = list(_generic_table_style(table_style, font_names["bold"], font_names["normal"],
                                       zebra=False).getCommands())
    cmds.append(("BACKGROUND", (0, 1), (-1, 1),
                 colors.HexColor("#EEEEEE") if table_style == "formal" else colors.HexColor("#e8f1ff")))
    t.setStyle(TableStyle(cmds))
    return t


FIG_DPI = 300  # calidad de impresión/publicación para las tres salidas (PDF/Word/Excel)


def _fig_to_png_bytes(fig, dpi=FIG_DPI) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight")
    buf.seek(0)
    return buf.read()


def _fig_to_image(fig, width=16 * cm):
    buf = io.BytesIO(_fig_to_png_bytes(fig))
    h = width * fig.get_size_inches()[1] / fig.get_size_inches()[0]
    return Image(buf, width=width, height=h)


def build_results_pdf(out_path: str, dataset_name: str, stats_table: pd.DataFrame,
                       figures: list[tuple[str, "object"]], params: dict | None = None,
                       ranking: tuple[pd.DataFrame, list[str]] | None = None,
                       multi_ranking: tuple[pd.DataFrame, list[str]] | None = None,
                       calc_detail: dict[str, dict] | None = None,
                       style: dict | None = None):
    """Genera un PDF con la tabla de indicadores estadísticos y una
    lista de figuras (título, Figure de matplotlib).

    Parameters
    ----------
    out_path : ruta del PDF a generar.
    dataset_name : nombre descriptivo del archivo/serie analizada.
    stats_table : DataFrame con los indicadores (una fila por corrida
        simulada o por indicador).
    figures : lista de tuplas (titulo, matplotlib.figure.Figure).
    params : diccionario opcional con los parámetros usados en el
        análisis (umbral POT, separación mínima, lambda Box-Cox...).
    ranking : tupla opcional (tabla_de_ranking, indicadores_usados),
        tal como la devuelve `ct_tool.stats.rank_runs`. Si se pasa, se
        agrega una sección con el puntaje compuesto y la corrida que
        mejor se ajusta.
    multi_ranking : tupla opcional (tabla_de_ranking, indicadores_usados)
        tal como la devuelve `ct_tool.stats.average_ranking_by_model`
        (del módulo de análisis múltiple/en paralelo). Si se pasa, se
        agrega una sección con el puntaje promedio de cada modelo a
        través de todas las variables analizadas.
    calc_detail : diccionario opcional {nombre_corrida: detalle}, tal
        como lo arma la interfaz al ejecutar el análisis (ver
        `ct_tool.stats.goodness_of_fit_detailed`). Si se pasa, cada
        sección de indicador incluye además "el paso del cálculo": con
        qué datos se trabajó, los cálculos intermedios reales y el
        resultado, no solo el número final.
    style : diccionario opcional {"font_family", "paragraph_align",
        "table_style"} (ver `DEFAULT_REPORT_STYLE`). Controla la
        tipografía de todo el PDF, la alineación de los párrafos de texto
        corrido y el aspecto de las tablas ("moderno" o "formal").
    """
    rstyle = _normalize_style(style)
    fonts = _pdf_font_names(rstyle["font_family"])
    align = _PDF_ALIGN.get(rstyle["paragraph_align"], TA_LEFT)
    table_style_name = rstyle["table_style"]

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleCT", parent=styles["Title"], fontSize=18,
                                  fontName=fonts["bold"])
    h2 = ParagraphStyle("H2CT", parent=styles["Heading2"], spaceBefore=10,
                         fontName=fonts["bold"])
    normal = ParagraphStyle("NormalCT", parent=styles["Normal"], fontName=fonts["normal"],
                             alignment=align)
    small = ParagraphStyle("SmallCT", parent=styles["Normal"], fontSize=7.5, leading=9.5,
                            fontName=fonts["normal"], alignment=align)

    doc = SimpleDocTemplate(out_path, pagesize=A4,
                             leftMargin=1.8 * cm, rightMargin=1.8 * cm,
                             topMargin=1.8 * cm, bottomMargin=1.8 * cm)
    story = []

    story.append(Paragraph("Reporte de análisis estadístico - CT (Comparison Tool)", title_style))
    story.append(Paragraph(f"Conjunto de datos: <b>{dataset_name}</b>", normal))
    story.append(Paragraph(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}", normal))
    story.append(Paragraph(f"<i>{AUTHOR_LINE}</i>", small))
    story.append(Spacer(1, 0.4 * cm))

    story.append(Paragraph("Metodología", h2))
    story.append(Paragraph(METHODOLOGY_INTRO, normal))
    story.append(Paragraph(METHODOLOGY_INDICADORES, normal))
    story.append(Paragraph(METHODOLOGY_AUSCULTACION, normal))
    story.append(Paragraph(METHODOLOGY_EXTREMOS, normal))
    story.append(Spacer(1, 0.3 * cm))

    if params:
        story.append(Paragraph("Parámetros del análisis", h2))
        rows = [["Parámetro", "Valor"]] + [[k, str(v)] for k, v in params.items()]
        t = Table(rows, hAlign="LEFT", colWidths=[7 * cm, 8 * cm])
        t.setStyle(_generic_table_style(table_style_name, fonts["bold"], fonts["normal"]))
        story.append(t)
        story.append(Spacer(1, 0.4 * cm))

    story.append(Paragraph("Indicadores de bondad de ajuste", h2))
    df = stats_table.reset_index() if stats_table.index.name else stats_table.copy()
    header = [str(c) for c in df.columns]
    story.append(_build_summary_table(df, header, font_names=fonts, table_style=table_style_name))

    story.append(Spacer(1, 0.3 * cm))
    story.append(Paragraph("Cómo se calcula cada indicador — teoría, fórmula, "
                            "procedimiento y resultados", h2))
    story.append(Paragraph(
        "No solo el resultado final: para cada método de validación se detalla qué mide, la "
        "fórmula exacta, el procedimiento de cálculo paso a paso, la escala de interpretación "
        "(Bueno / Regular / Malo) y el resultado obtenido en esta corrida para cada serie "
        "simulada.", normal))
    story.append(Spacer(1, 0.15 * cm))

    h3 = ParagraphStyle("H3CT", parent=styles["Heading3"], fontSize=11.5, fontName=fonts["bold"],
                         textColor=colors.HexColor("#2C6FD1"), spaceBefore=10, spaceAfter=2)
    # La "caja" de fórmula se mantiene siempre en Courier (monoespaciada):
    # es la que da el aspecto de fórmula/código, independiente de la
    # tipografía elegida para el resto del texto corrido.
    mono = ParagraphStyle("MonoCT", parent=styles["Normal"], fontName="Courier",
                           fontSize=8.5, leading=11, backColor=colors.HexColor("#F1F3F8"),
                           borderPadding=(4, 6, 4, 6))
    body_small = ParagraphStyle("BodySmallCT", parent=styles["Normal"], fontName=fonts["normal"],
                                 fontSize=8.5, leading=11, alignment=align)

    for key in header:
        help_entry = INDICATOR_HELP.get(key)
        if not help_entry:
            continue
        story.append(Paragraph(key, h3))
        story.append(Paragraph(f"<b>¿Qué es y para qué sirve?</b> {help_entry.get('teoria', '')}",
                                body_small))
        story.append(Spacer(1, 0.1 * cm))
        story.append(Paragraph(help_entry["formula"], mono))
        pasos = help_entry.get("pasos", [])
        if pasos:
            steps_text = "<br/>".join(f"{i}. {p}" for i, p in enumerate(pasos, start=1))
            story.append(Paragraph(f"<b>¿Cómo se determina el resultado?</b><br/>{steps_text}",
                                    body_small))
        story.append(Paragraph(f"<b>Escala de interpretación:</b> {scale_description(key)}",
                                body_small))
        rows = _indicator_results_rows(df, key)
        if rows:
            table_rows = [["Corrida", "Valor", "Calificación"]]
            style_cmds = list(_generic_table_style(table_style_name, fonts["bold"], fonts["normal"],
                                                     zebra=False).getCommands())
            for i, (run_name, val, label) in enumerate(rows, start=1):
                table_rows.append([run_name, f"{val:.4f}" if isinstance(val, (int, float)) else str(val),
                                    label or "—"])
                bg = _QUALITY_BG_HEX.get(label)
                if bg:
                    style_cmds.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor(f"#{bg}")))
            tres = Table(table_rows, hAlign="LEFT", colWidths=[6 * cm, 3 * cm, 3.5 * cm])
            tres.setStyle(TableStyle(style_cmds))
            story.append(tres)
        trace_lines = _calc_trace_lines(df, key, calc_detail)
        if trace_lines:
            story.append(Spacer(1, 0.1 * cm))
            story.append(Paragraph("<b>Cálculo con los datos de esta corrida (paso a paso):</b>",
                                    body_small))
            trace_style = ParagraphStyle("TraceCT", parent=body_small, fontSize=7.3, leading=9.5,
                                          textColor=colors.HexColor("#334155"))
            for line in trace_lines:
                story.append(Paragraph(line, trace_style))
            story.append(Spacer(1, 0.1 * cm))
        gauge_fig = _indicator_gauge_figure(df, key)
        if gauge_fig is not None:
            story.append(_fig_to_image(gauge_fig, width=10 * cm))
        story.append(Spacer(1, 0.3 * cm))

    if ranking is not None:
        ranking_df, used_indicators = ranking
        story.append(Spacer(1, 0.4 * cm))
        story.append(Paragraph("Ranking compuesto y mejor ajuste", h2))
        story.append(Paragraph(_ranking_verdict_text(ranking_df, used_indicators), normal))
        story.append(Spacer(1, 0.2 * cm))
        rk = ranking_df[["Puntaje compuesto (0-1)"]].reset_index()
        rk.columns = ["Corrida", "Puntaje compuesto (0-1)"]
        rk_rows = [list(rk.columns)] + rk.round(4).astype(str).values.tolist()
        trk = Table(rk_rows, hAlign="LEFT", colWidths=[8 * cm, 5 * cm])
        trk_cmds = list(_generic_table_style(table_style_name, fonts["bold"], fonts["normal"],
                                               zebra=False).getCommands())
        trk_cmds.append(("BACKGROUND", (0, 1), (-1, 1),
                          colors.HexColor("#EEEEEE") if table_style_name == "formal"
                          else colors.HexColor("#e8f1ff")))
        trk.setStyle(TableStyle(trk_cmds))
        story.append(trk)

    if figures:
        story.append(PageBreak())
        story.append(Paragraph("Gráficos de diagnóstico", h2))
        story.append(Paragraph(
            "Cada gráfico incluye a continuación una breve explicación de qué muestra y cómo "
            "interpretarlo, no solo la imagen.", normal))
        story.append(Spacer(1, 0.15 * cm))
    # Cada figura (título + descripción + imagen) va en un único bloque
    # `KeepTogether`, para que nunca quede partida entre dos páginas --
    # pero SIN forzar un salto de página antes de cada una: reportlab
    # solo pasa a la siguiente página cuando el bloque no entra en el
    # espacio que queda, así que varias figuras que sí entran juntas en
    # una misma hoja quedan una junto a otra, en vez de dejar el resto
    # de la página anterior en blanco (que es lo que pasaba antes, con
    # un `PageBreak()` fijo delante de cada gráfico).
    for title, fig in figures:
        block = [Paragraph(title, h2)]
        caption = FIGURE_DESCRIPTIONS.get(title)
        if caption:
            block.append(Paragraph(caption, normal))
            block.append(Spacer(1, 0.15 * cm))
        block.append(_fig_to_image(fig))
        block.append(Spacer(1, 0.3 * cm))
        story.append(KeepTogether(block))

    if multi_ranking is not None:
        mr_df, mr_used = multi_ranking
        story.append(Spacer(1, 0.3 * cm))
        story.append(Paragraph("Ranking promedio por modelo (análisis múltiple)", h2))
        story.append(Paragraph(_multi_ranking_text(mr_df, mr_used), normal))
        story.append(Spacer(1, 0.15 * cm))
        story.append(_build_multi_ranking_table(mr_df, font_names=fonts, table_style=table_style_name))

    doc.build(story, onFirstPage=_draw_pdf_watermark, onLaterPages=_draw_pdf_watermark)
    return out_path


def _draw_pdf_watermark(canvas, doc):
    """Marca de agua por defecto en cada página del PDF: el logo de
    Villazón, centrado y semi-transparente, para identificar el
    origen del reporte sin estorbar la lectura. Si el archivo del
    logo no está disponible por algún motivo (p. ej. un empaquetado
    incompleto), se omite en silencio -- la marca de agua es un
    detalle estético, nunca debe impedir que el reporte se genere."""
    if not WATERMARK_IMAGE_PATH.exists():
        return
    canvas.saveState()
    try:
        canvas.setFillAlpha(0.10)
    except AttributeError:
        pass  # versiones muy antiguas de reportlab sin soporte de alpha
    try:
        img = ImageReader(str(WATERMARK_IMAGE_PATH))
        iw, ih = img.getSize()
        page_w, page_h = A4
        # Tamaño moderado, centrado en la página, sin rotar -- el logo
        # ya trae su propio texto, y un logo rotado se lee peor que
        # una marca de agua de solo texto.
        draw_w = page_w * 0.55
        draw_h = draw_w * (ih / iw)
        canvas.drawImage(img, (page_w - draw_w) / 2, (page_h - draw_h) / 2,
                          width=draw_w, height=draw_h, mask="auto",
                          preserveAspectRatio=True)
    except Exception:
        pass
    canvas.restoreState()


# ---------------------------------------------------------------------------
# Reporte Word (.docx)
# ---------------------------------------------------------------------------

def _docx_set_cell_shading(cell, hex_color: str):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _docx_set_cell_border(cell, edge: str, size: int = 8, color: str = "000000"):
    """Agrega un filete a un solo lado ("top"/"bottom") de una celda de
    Word, vía manipulación directa del XML (`python-docx` no tiene una
    API de alto nivel para bordes de celda). `size` está en octavos de
    punto (igual que la API de Word), p. ej. 16 ≈ 2pt."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    tag = qn(f"w:{edge}")
    el = tcBorders.find(tag)
    if el is None:
        el = OxmlElement(f"w:{edge}")
        tcBorders.append(el)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:color"), color)


def _docx_apply_booktabs_rules(table):
    """Deja SOLO tres filetes horizontales en una tabla Word -- uno
    arriba del encabezado, uno bajo el encabezado y uno al pie de la
    tabla -- sin grilla ni líneas verticales ni entre filas de datos,
    igual que la variante "formal" del PDF/Excel (ver
    `report._generic_table_style`). Se usa en vez de `table.style =
    "Table Grid"` cuando el estilo elegido es "formal"."""
    header_cells = table.rows[0].cells
    for cell in header_cells:
        _docx_set_cell_border(cell, "top", size=16)
        _docx_set_cell_border(cell, "bottom", size=10)
    for cell in table.rows[-1].cells:
        _docx_set_cell_border(cell, "bottom", size=16)


def _docx_set_table_grid(table, table_style: str):
    """"moderno": cuadrícula completa (`"Table Grid"`, aspecto original).
    "formal": sin estilo de cuadrícula -- los tres filetes se agregan
    aparte con `_docx_apply_booktabs_rules` una vez que la tabla ya
    tiene todas sus filas (ver cada sitio donde se llama)."""
    if table_style != "formal":
        table.style = "Table Grid"


def _docx_header_row(table, values, font_name="Arial", table_style="moderno"):
    from docx.shared import Pt, RGBColor
    formal = table_style == "formal"
    row = table.rows[0]
    for cell, val in zip(row.cells, values):
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(str(val))
        run.bold = True
        run.font.name = font_name
        run.font.size = Pt(9)
        if formal:
            run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
            _docx_set_cell_shading(cell, "FFFFFF")
        else:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _docx_set_cell_shading(cell, NAVY_HEX)


def _docx_data_rows(table, rows, font_name="Arial", table_style="moderno", start=1):
    from docx.shared import Pt
    zebra = table_style != "formal"
    for r, values in enumerate(rows, start=start):
        for cell, val in zip(table.rows[r].cells, values):
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            run.font.name = font_name
            run.font.size = Pt(9)
        if zebra and (r - start) % 2 == 1:
            for cell in table.rows[r].cells:
                _docx_set_cell_shading(cell, "F2F5F8")


def build_results_docx(out_path: str, dataset_name: str, stats_table: pd.DataFrame,
                        figures: list[tuple[str, "object"]], params: dict | None = None,
                        ranking: tuple[pd.DataFrame, list[str]] | None = None,
                        multi_ranking: tuple[pd.DataFrame, list[str]] | None = None,
                        calc_detail: dict[str, dict] | None = None,
                        style: dict | None = None):
    """Genera un reporte Word (.docx) con la misma tabla de indicadores
    y figuras que `build_results_pdf`, en formato editable por el
    usuario (por ejemplo para pegarlo en su tesis o informe).

    `ranking` (opcional): tupla (tabla_de_ranking, indicadores_usados)
    de `ct_tool.stats.rank_runs`, para agregar la sección de puntaje
    compuesto y mejor ajuste.

    `multi_ranking` (opcional): tupla (tabla_de_ranking, indicadores_usados)
    de `ct_tool.stats.average_ranking_by_model` (análisis múltiple/en
    paralelo), para agregar la sección de puntaje promedio por modelo.

    `style` (opcional): mismo diccionario {"font_family", "paragraph_align",
    "table_style"} que recibe `build_results_pdf` (ver `DEFAULT_REPORT_STYLE`)."""
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT, WD_SECTION

    rstyle = _normalize_style(style)
    font_name = rstyle["font_family"]
    table_style_name = rstyle["table_style"]
    para_align = {"left": WD_ALIGN_PARAGRAPH.LEFT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                  "center": WD_ALIGN_PARAGRAPH.CENTER}.get(rstyle["paragraph_align"],
                                                             WD_ALIGN_PARAGRAPH.LEFT)

    MARGIN_CM = 1.8

    def _set_margins(section, margin_cm=MARGIN_CM):
        section.left_margin = Cm(margin_cm)
        section.right_margin = Cm(margin_cm)
        section.top_margin = Cm(margin_cm)
        section.bottom_margin = Cm(margin_cm)

    def _content_width_cm(section):
        return (section.page_width - section.left_margin - section.right_margin) / 360000

    def _body_p(text: str = ""):
        """Párrafo de texto corrido con la alineación elegida por la
        usuaria (justificado/izquierda/centrado) -- se usa para todo el
        texto explicativo (metodología, teoría de cada indicador, pasos,
        escala...), no para encabezados ni celdas de tabla."""
        p = doc.add_paragraph(text)
        p.alignment = para_align
        return p

    doc = Document()
    doc.styles["Normal"].font.name = font_name
    doc.styles["Normal"].font.size = Pt(10)
    _set_margins(doc.sections[0])

    title = doc.add_heading(level=0)
    title_run = title.add_run("Reporte de análisis estadístico — CT (Comparison Tool)")
    title_run.font.name = font_name
    title_run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
    title_run.font.size = Pt(20)

    p = doc.add_paragraph()
    p.add_run("Conjunto de datos: ").bold = False
    p.add_run(str(dataset_name)).bold = True
    doc.add_paragraph(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    p_author = doc.add_paragraph()
    run_author = p_author.add_run(AUTHOR_LINE)
    run_author.italic = True
    run_author.font.size = Pt(8.5)

    doc.add_heading("Metodología", level=1)
    _body_p(METHODOLOGY_INTRO)
    _body_p(METHODOLOGY_INDICADORES)
    _body_p(METHODOLOGY_AUSCULTACION)
    _body_p(METHODOLOGY_EXTREMOS)

    if params:
        doc.add_heading("Parámetros del análisis", level=1)
        t = doc.add_table(rows=1 + len(params), cols=2)
        _docx_set_table_grid(t, table_style_name)
        t.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _docx_header_row(t, ["Parámetro", "Valor"], font_name, table_style_name)
        _docx_data_rows(t, [[k, str(v)] for k, v in params.items()], font_name, table_style_name)
        if table_style_name == "formal":
            _docx_apply_booktabs_rules(t)

    # La tabla de indicadores suele tener muchas columnas (n, MAE, RMSE,
    # EF, CD, R, ...): en vertical (A4) no entra sin que el encabezado se
    # parta en varias líneas por columna. Si hay más de 6 columnas, se usa
    # una página en horizontal solo para esta tabla, y se vuelve a
    # vertical para el resto del reporte (parámetros ya quedaron arriba).
    header = [str(c) for c in (stats_table.reset_index() if stats_table.index.name
                                else stats_table.copy()).columns]
    df = stats_table.reset_index() if stats_table.index.name else stats_table.copy()
    body = df.round(4).astype(str).values.tolist()
    wide_table = len(header) > 6

    if wide_table:
        landscape_section = doc.add_section(WD_SECTION.NEW_PAGE)
        landscape_section.orientation = WD_ORIENT.LANDSCAPE
        landscape_section.page_width, landscape_section.page_height = (
            landscape_section.page_height, landscape_section.page_width)
        _set_margins(landscape_section)
        doc.add_heading("Indicadores de bondad de ajuste", level=1)
    else:
        doc.add_heading("Indicadores de bondad de ajuste", level=1)

    t = doc.add_table(rows=1 + len(body), cols=len(header))
    _docx_set_table_grid(t, table_style_name)
    t.autofit = False
    avail_cm = _content_width_cm(doc.sections[-1])
    col_w = Cm(max(1.4, avail_cm / len(header)))
    for col in t.columns:
        col.width = col_w
    _docx_header_row(t, header, font_name, table_style_name)
    _docx_data_rows(t, body, font_name, table_style_name)
    if table_style_name == "formal":
        _docx_apply_booktabs_rules(t)

    if wide_table:
        portrait_section = doc.add_section(WD_SECTION.NEW_PAGE)
        portrait_section.orientation = WD_ORIENT.PORTRAIT
        portrait_section.page_width, portrait_section.page_height = (
            portrait_section.page_height, portrait_section.page_width)
        _set_margins(portrait_section)

    content_width_cm = _content_width_cm(doc.sections[-1])

    doc.add_heading("Cómo se calcula cada indicador — teoría, fórmula, procedimiento y resultados",
                     level=1)
    _body_p(
        "No solo el resultado final: para cada método de validación se detalla qué mide, la "
        "fórmula exacta, el procedimiento de cálculo paso a paso, la escala de interpretación "
        "(Bueno / Regular / Malo) y el resultado obtenido en esta corrida para cada serie "
        "simulada.")
    for key in header:
        help_entry = INDICATOR_HELP.get(key)
        if not help_entry:
            continue
        doc.add_heading(key, level=2)

        p_theory = _body_p()
        p_theory.add_run("¿Qué es y para qué sirve? ").bold = True
        p_theory.add_run(help_entry.get("teoria", ""))

        p_formula = doc.add_paragraph()
        run_f = p_formula.add_run(help_entry["formula"])
        run_f.font.name = "Consolas"
        run_f.font.size = Pt(10)

        pasos = help_entry.get("pasos", [])
        if pasos:
            p_steps_label = _body_p()
            p_steps_label.add_run("¿Cómo se determina el resultado?").bold = True
            # Numeración manual (texto "1. ", "2. "...) en vez del estilo
            # "List Number" de Word: ese estilo comparte una única lista
            # numerada para todo el documento, así que los pasos de cada
            # indicador seguirían la numeración del anterior (p. ej. STD
            # empezaría en "5." en vez de "1.") en vez de reiniciar en 1
            # para cada indicador.
            for i, paso in enumerate(pasos, start=1):
                p_step = _body_p()
                p_step.paragraph_format.left_indent = Pt(14)
                p_step.add_run(f"{i}. {paso}")

        p_scale = _body_p()
        p_scale.add_run("Escala de interpretación: ").bold = True
        p_scale.add_run(scale_description(key))

        rows = _indicator_results_rows(df, key)
        if rows:
            tind = doc.add_table(rows=1 + len(rows), cols=3)
            _docx_set_table_grid(tind, table_style_name)
            _docx_header_row(tind, ["Corrida", "Valor", "Calificación"], font_name, table_style_name)
            for r, (run_name, val, label) in enumerate(rows, start=1):
                val_text = f"{val:.4f}" if isinstance(val, (int, float)) else str(val)
                for cell, text in zip(tind.rows[r].cells, [run_name, val_text, label or "—"]):
                    cell.text = ""
                    run_c = cell.paragraphs[0].add_run(text)
                    run_c.font.size = Pt(9)
                bg = _QUALITY_BG_HEX.get(label)
                if bg:
                    for cell in tind.rows[r].cells:
                        _docx_set_cell_shading(cell, bg)
            if table_style_name == "formal":
                _docx_apply_booktabs_rules(tind)
        trace_lines = _calc_trace_lines(df, key, calc_detail)
        if trace_lines:
            p_trace_label = _body_p()
            p_trace_label.add_run("Cálculo con los datos de esta corrida (paso a paso):").bold = True
            for line in trace_lines:
                p_trace = _body_p()
                p_trace.paragraph_format.left_indent = Pt(10)
                # `line` viene con un único tag <b> al inicio (nombre de
                # la corrida) seguido de texto plano -- se separa a mano
                # en vez de meter un parser HTML completo solo para esto.
                plain = (line.replace("<b>", "").replace("</b>", "")
                             .replace("&nbsp;", " "))
                run_t = p_trace.add_run(plain)
                run_t.font.size = Pt(8)
        gauge_fig = _indicator_gauge_figure(df, key)
        if gauge_fig is not None:
            gauge_png = _fig_to_png_bytes(gauge_fig)
            gauge_para = doc.add_paragraph()
            gauge_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            gauge_para.add_run().add_picture(io.BytesIO(gauge_png), width=Cm(min(content_width_cm, 11.0)))
        doc.add_paragraph()

    if ranking is not None:
        ranking_df, used_indicators = ranking
        doc.add_heading("Ranking compuesto y mejor ajuste", level=1)
        _body_p(_ranking_verdict_text(ranking_df, used_indicators))
        rk = ranking_df[["Puntaje compuesto (0-1)"]].reset_index()
        rk.columns = ["Corrida", "Puntaje compuesto (0-1)"]
        trk = doc.add_table(rows=1 + len(rk), cols=2)
        _docx_set_table_grid(trk, table_style_name)
        _docx_header_row(trk, list(rk.columns), font_name, table_style_name)
        _docx_data_rows(trk, rk.round(4).astype(str).values.tolist(), font_name, table_style_name)
        if table_style_name == "formal":
            _docx_apply_booktabs_rules(trk)

    if figures:
        doc.add_page_break()
        doc.add_heading("Gráficos de diagnóstico", level=1)
        _body_p("Cada gráfico incluye a continuación una breve explicación de qué "
                "muestra y cómo interpretarlo, no solo la imagen.")
    # Sin salto de página forzado antes de CADA gráfico (a diferencia de
    # antes): Word fluye el contenido de forma continua, así que varios
    # gráficos que entren en la misma hoja quedan uno junto al otro, en
    # vez de dejar el resto de la página anterior en blanco. El título
    # se marca "mantener con el siguiente" para que, si Word sí necesita
    # cortar página, no separe el título de su imagen.
    for fig_title, fig in figures:
        heading = doc.add_heading(fig_title, level=1)
        heading.paragraph_format.keep_with_next = True
        caption = FIGURE_DESCRIPTIONS.get(fig_title)
        if caption:
            cap_para = _body_p(caption)
            cap_para.paragraph_format.keep_with_next = True
        png_bytes = _fig_to_png_bytes(fig)
        w_in, h_in = fig.get_size_inches()
        pic_width_cm = min(content_width_cm, 16.0)
        pic_para = doc.add_paragraph()
        pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = pic_para.add_run()
        run.add_picture(io.BytesIO(png_bytes), width=Cm(pic_width_cm))

    if multi_ranking is not None:
        mr_df, mr_used = multi_ranking
        doc.add_heading("Ranking promedio por modelo (análisis múltiple)", level=1)
        _body_p(_multi_ranking_text(mr_df, mr_used))
        mr_cols = [c for c in ["Puntaje promedio (0-1)", "N° de variables"] if c in mr_df.columns]
        mrk = mr_df[mr_cols].reset_index()
        mrk.columns = ["Modelo"] + mr_cols
        tmr = doc.add_table(rows=1 + len(mrk), cols=len(mrk.columns))
        _docx_set_table_grid(tmr, table_style_name)
        _docx_header_row(tmr, list(mrk.columns), font_name, table_style_name)
        _docx_data_rows(tmr, mrk.round(4).astype(str).values.tolist(), font_name, table_style_name)
        if table_style_name == "formal":
            _docx_apply_booktabs_rules(tmr)

    doc.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
# Reporte Excel (.xlsx)
# ---------------------------------------------------------------------------

def build_results_xlsx(out_path: str, dataset_name: str, stats_table: pd.DataFrame,
                        figures: list[tuple[str, "object"]], params: dict | None = None,
                        ranking: tuple[pd.DataFrame, list[str]] | None = None,
                        multi_ranking: tuple[pd.DataFrame, list[str]] | None = None,
                        calc_detail: dict[str, dict] | None = None,
                        style: dict | None = None):
    """Genera un reporte Excel (.xlsx) con una hoja "Resumen" (metodología,
    parámetros, tabla de indicadores, fórmulas y ranking compuesto) y una
    hoja por cada figura, con la imagen incrustada. Es un reporte de la
    corrida (como el PDF/Word), con valores ya calculados — no una
    planilla de fórmulas recalculables, porque resume series de tiempo
    completas y un análisis de extremos que no tendría sentido
    reconstruir celda por celda en Excel.

    `ranking` (opcional): tupla (tabla_de_ranking, indicadores_usados)
    de `ct_tool.stats.rank_runs`.

    `multi_ranking` (opcional): tupla (tabla_de_ranking, indicadores_usados)
    de `ct_tool.stats.average_ranking_by_model` (análisis múltiple/en
    paralelo).

    `style` (opcional): mismo diccionario {"font_family", "paragraph_align",
    "table_style"} que recibe `build_results_pdf` (ver `DEFAULT_REPORT_STYLE`)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.image import Image as XLImage

    rstyle = _normalize_style(style)
    formal = rstyle["table_style"] == "formal"
    text_align = "justify" if rstyle["paragraph_align"] == "justify" else (
        "center" if rstyle["paragraph_align"] == "center" else "left")

    FONT_NAME = rstyle["font_family"]
    NAVY = "2C3E50"
    LIGHT = "F2F5F8"
    HEAD_FONT = (Font(name=FONT_NAME, bold=True, color="000000", size=10) if formal
                 else Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10))
    HEAD_FILL = (PatternFill("solid", fgColor="FFFFFF") if formal
                 else PatternFill("solid", fgColor=NAVY))
    LIGHT_FILL = PatternFill("solid", fgColor=LIGHT)
    BASE_FONT = Font(name=FONT_NAME, size=10)
    TITLE_FONT = Font(name=FONT_NAME, bold=True, size=16, color=NAVY)
    THIN = Side(style="thin", color="BBBBBB")
    THICK = Side(style="medium", color="000000")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    NO_BORDER = Border()
    BOTTOM_RULE = Border(bottom=THICK)
    # "formal": sin cuadrícula -- solo tres filetes horizontales (arriba
    # y debajo del encabezado, y al pie de cada tabla), igual que el
    # PDF/Word (ver `report._generic_table_style`). `DATA_BORDER` se usa
    # en las filas de datos; el filete del pie se agrega aparte, encima
    # de la última fila ya escrita de cada tabla (ver cada sitio donde
    # se usa `BOTTOM_RULE`).
    HEAD_BORDER = (Border(top=THICK, bottom=THICK) if formal else BORDER)
    DATA_BORDER = (NO_BORDER if formal else BORDER)

    def _zebra(row_cell, i):
        """Sombreado alternado solo en el estilo "moderno" -- el "formal"
        (tablas académicas) va siempre en blanco, sin filas alternadas."""
        if not formal and i % 2 == 1:
            row_cell.fill = LIGHT_FILL

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"

    ws.cell(row=1, column=1, value="Reporte de análisis estadístico - CT (Comparison Tool)").font = TITLE_FONT
    ws.cell(row=2, column=1, value=f"Conjunto de datos: {dataset_name}").font = BASE_FONT
    ws.cell(row=3, column=1, value=f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}").font = BASE_FONT
    ws.cell(row=4, column=1, value=AUTHOR_LINE).font = Font(name=FONT_NAME, italic=True, size=8.5)

    def _write_paragraph_block(row_start, text, span_cols=9, span_rows=4):
        ws.merge_cells(start_row=row_start, start_column=1, end_row=row_start + span_rows - 1,
                        end_column=span_cols)
        c = ws.cell(row=row_start, column=1, value=text)
        c.font = BASE_FONT
        c.alignment = Alignment(horizontal=text_align, vertical="top", wrap_text=True)
        return row_start + span_rows

    row = 5
    ws.cell(row=row, column=1, value="Metodología").font = Font(name=FONT_NAME, bold=True, size=12)
    row += 1
    row = _write_paragraph_block(row, METHODOLOGY_INTRO)
    row = _write_paragraph_block(row, METHODOLOGY_INDICADORES)
    row = _write_paragraph_block(row, METHODOLOGY_AUSCULTACION)
    row = _write_paragraph_block(row, METHODOLOGY_EXTREMOS)
    row += 1

    if params:
        ws.cell(row=row, column=1, value="Parámetros del análisis").font = Font(name=FONT_NAME, bold=True, size=12)
        row += 1
        for j, h in enumerate(["Parámetro", "Valor"], start=1):
            c = ws.cell(row=row, column=j, value=h)
            c.font = HEAD_FONT; c.fill = HEAD_FILL; c.border = HEAD_BORDER
        row += 1
        for k, v in params.items():
            ws.cell(row=row, column=1, value=str(k)).font = BASE_FONT
            ws.cell(row=row, column=2, value=str(v)).font = BASE_FONT
            for col in (1, 2):
                ws.cell(row=row, column=col).border = DATA_BORDER
            row += 1
        if formal and params:
            for col in (1, 2):
                ws.cell(row=row - 1, column=col).border = BOTTOM_RULE
        row += 1

    ws.cell(row=row, column=1, value="Indicadores de bondad de ajuste").font = Font(name=FONT_NAME, bold=True, size=12)
    row += 1
    df = stats_table.reset_index() if stats_table.index.name else stats_table.copy()
    header = [str(c) for c in df.columns]
    header_row = row
    for j, h in enumerate(header, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = HEAD_FONT; c.fill = HEAD_FILL; c.border = HEAD_BORDER
        c.alignment = Alignment(horizontal="center")
    row += 1
    for i, (_, r) in enumerate(df.iterrows()):
        for j, col in enumerate(header, start=1):
            val = r[col]
            val = round(val, 4) if isinstance(val, float) else val
            c = ws.cell(row=row, column=j, value=val)
            c.font = BASE_FONT; c.border = DATA_BORDER
            if not formal and i % 2 == 1:
                c.fill = LIGHT_FILL
        row += 1
    if formal and len(df):
        for j in range(1, len(header) + 1):
            ws.cell(row=row - 1, column=j).border = BOTTOM_RULE
    for j, h in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(22, 3 + len(h)))
    if "Corrida" in header:
        longest_run_name = max((len(str(v)) for v in df["Corrida"]), default=0)
        ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0,
                                               min(28, longest_run_name + 3))
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate
    row += 1

    ws.cell(row=row, column=1,
            value="Cómo se calcula cada indicador — teoría, fórmula, procedimiento y resultados"
            ).font = Font(name=FONT_NAME, bold=True, size=12)
    row += 1
    row = _write_paragraph_block(
        row,
        "No solo el resultado final: para cada método de validación se detalla qué mide, la "
        "fórmula exacta, el procedimiento de cálculo paso a paso, la escala de interpretación "
        "(Bueno / Regular / Malo) y el resultado obtenido en esta corrida para cada serie "
        "simulada.", span_rows=2)
    row += 1

    formula_entries = [(key, INDICATOR_HELP[key]) for key in header if key in INDICATOR_HELP]
    for key, help_entry in formula_entries:
        ws.cell(row=row, column=1, value=key).font = Font(name=FONT_NAME, bold=True, size=12, color=NAVY)
        row += 1
        row = _write_paragraph_block(
            row, f"Qué es y para qué sirve: {help_entry.get('teoria', '')}", span_rows=4)

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        c_formula = ws.cell(row=row, column=1, value=help_entry["formula"])
        c_formula.font = Font(name="Consolas", size=10)
        c_formula.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        row += 1

        pasos = help_entry.get("pasos", [])
        if pasos:
            steps_text = "¿Cómo se determina el resultado?\n" + "\n".join(
                f"{i}. {p}" for i, p in enumerate(pasos, start=1))
            row = _write_paragraph_block(row, steps_text, span_rows=max(3, len(pasos) + 1))

        row = _write_paragraph_block(row, f"Escala de interpretación: {scale_description(key)}",
                                      span_rows=2)

        rows_data = _indicator_results_rows(df, key)
        if rows_data:
            for j, h in enumerate(["Corrida", "Valor", "Calificación"], start=1):
                c = ws.cell(row=row, column=j, value=h)
                c.font = HEAD_FONT; c.fill = HEAD_FILL; c.border = HEAD_BORDER
            row += 1
            for run_name, val, label in rows_data:
                val_text = round(val, 4) if isinstance(val, float) else val
                ws.cell(row=row, column=1, value=run_name).font = BASE_FONT
                ws.cell(row=row, column=2, value=val_text).font = BASE_FONT
                ws.cell(row=row, column=3, value=label or "—").font = BASE_FONT
                bg = _QUALITY_BG_HEX.get(label)
                if bg:
                    for col in (1, 2, 3):
                        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=bg)
                for col in (1, 2, 3):
                    ws.cell(row=row, column=col).border = DATA_BORDER
                row += 1
            if formal:
                for col in (1, 2, 3):
                    ws.cell(row=row - 1, column=col).border = BOTTOM_RULE
        trace_lines = _calc_trace_lines(df, key, calc_detail)
        if trace_lines:
            plain_lines = [
                l.replace("<b>", "").replace("</b>", "").replace("&nbsp;", " ")
                for l in trace_lines]
            row = _write_paragraph_block(
                row, "Cálculo con los datos de esta corrida (paso a paso):\n" +
                "\n".join(plain_lines), span_rows=max(2, len(plain_lines) + 1))
        gauge_fig = _indicator_gauge_figure(df, key)
        if gauge_fig is not None:
            gauge_png = _fig_to_png_bytes(gauge_fig, dpi=140)
            gauge_img = XLImage(io.BytesIO(gauge_png))
            gw_in, gh_in = gauge_fig.get_size_inches()
            target_w_px = 480
            gauge_img.width = target_w_px
            gauge_img.height = int(target_w_px * gh_in / gw_in)
            ws.add_image(gauge_img, f"A{row}")
            # Una fila de Excel mide ~20px de alto por defecto -- se
            # reservan suficientes filas para que el encabezado del
            # próximo indicador no quede tapado por la imagen (las
            # imágenes de openpyxl "flotan" sobre las celdas, no
            # empujan el contenido siguiente automáticamente).
            row += max(10, round(gauge_img.height / 20) + 1)
        row += 1

    if ranking is not None:
        ranking_df, used_indicators = ranking
        ws.cell(row=row, column=1, value="Ranking compuesto y mejor ajuste").font = Font(name=FONT_NAME, bold=True, size=12)
        row += 1
        row = _write_paragraph_block(row, _ranking_verdict_text(ranking_df, used_indicators), span_rows=3)
        rk = ranking_df[["Puntaje compuesto (0-1)"]].reset_index()
        rk.columns = ["Corrida", "Puntaje compuesto (0-1)"]
        for j, h in enumerate(rk.columns, start=1):
            c = ws.cell(row=row, column=j, value=h)
            c.font = HEAD_FONT; c.fill = HEAD_FILL; c.border = HEAD_BORDER
        row += 1
        for i, (_, r) in enumerate(rk.iterrows()):
            ws.cell(row=row, column=1, value=str(r["Corrida"])).font = BASE_FONT
            ws.cell(row=row, column=2, value=round(float(r["Puntaje compuesto (0-1)"]), 4)).font = BASE_FONT
            if i == 0:
                for col in (1, 2):
                    ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=("EEEEEE" if formal else "E8F1FF"))
            for col in (1, 2):
                ws.cell(row=row, column=col).border = DATA_BORDER
            row += 1
        if formal and len(rk):
            for col in (1, 2):
                ws.cell(row=row - 1, column=col).border = BOTTOM_RULE
        row += 1

    if multi_ranking is not None:
        mr_df, mr_used = multi_ranking
        ws.cell(row=row, column=1, value="Ranking promedio por modelo (análisis múltiple)").font = Font(
            name=FONT_NAME, bold=True, size=12)
        row += 1
        row = _write_paragraph_block(row, _multi_ranking_text(mr_df, mr_used), span_rows=3)
        mr_cols = [c for c in ["Puntaje promedio (0-1)", "N° de variables"] if c in mr_df.columns]
        mrk = mr_df[mr_cols].reset_index()
        mrk.columns = ["Modelo"] + mr_cols
        for j, h in enumerate(mrk.columns, start=1):
            c = ws.cell(row=row, column=j, value=h)
            c.font = HEAD_FONT; c.fill = HEAD_FILL; c.border = HEAD_BORDER
        row += 1
        for i, (_, r) in enumerate(mrk.iterrows()):
            for j, col_name in enumerate(mrk.columns, start=1):
                val = r[col_name]
                val = round(float(val), 4) if col_name == "Puntaje promedio (0-1)" else val
                c = ws.cell(row=row, column=j, value=val)
                c.font = BASE_FONT; c.border = DATA_BORDER
                if i == 0:
                    c.fill = PatternFill("solid", fgColor=("EEEEEE" if formal else "E8F1FF"))
            row += 1
        if formal and len(mrk):
            for j in range(1, len(mrk.columns) + 1):
                ws.cell(row=row - 1, column=j).border = BOTTOM_RULE

    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{get_column_letter(max(len(header), 9))}{row - 1}"

    used_sheet_names = {"Resumen"}
    for fig_title, fig in figures:
        base = "".join(ch for ch in fig_title if ch not in '[]:*?/\\')[:28] or "Figura"
        name = base
        n = 1
        while name in used_sheet_names:
            n += 1
            name = f"{base[:28]}_{n}"
        used_sheet_names.add(name)
        wf = wb.create_sheet(name)
        wf.cell(row=1, column=1, value=fig_title).font = Font(name=FONT_NAME, bold=True, size=13, color=NAVY)
        caption = FIGURE_DESCRIPTIONS.get(fig_title)
        img_row = 3
        if caption:
            wf.merge_cells(start_row=2, start_column=1, end_row=2, end_column=9)
            c_cap = wf.cell(row=2, column=1, value=caption)
            c_cap.font = Font(name=FONT_NAME, size=9.5, italic=True)
            c_cap.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            wf.row_dimensions[2].height = 46
            img_row = 4
        png_bytes = _fig_to_png_bytes(fig, dpi=170)  # más liviano que 300dpi para no inflar el .xlsx
        img_stream = io.BytesIO(png_bytes)
        xl_img = XLImage(img_stream)
        w_in, h_in = fig.get_size_inches()
        target_w_px = 900
        xl_img.width = target_w_px
        xl_img.height = int(target_w_px * h_in / w_in)
        wf.add_image(xl_img, f"A{img_row}")
        wf.sheet_view.showGridLines = False
        wf.page_setup.orientation = "landscape"
        wf.page_setup.fitToWidth = 1
        wf.page_setup.fitToHeight = 1
        wf.sheet_properties.pageSetUpPr.fitToPage = True
        last_col = get_column_letter(max(2, round(target_w_px / 64)))
        last_row = img_row + max(15, round(xl_img.height / 20))
        wf.print_area = f"A1:{last_col}{last_row}"

    wb.save(out_path)
    return out_path
